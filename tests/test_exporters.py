import json
import sys
from pathlib import Path

import pytest

from app.export_manager import ExportError, check_directory, export_session, export_to_path
from app.exporters import EXPORTERS, ExportOptions
from app.exporters.base import EMPTY_NOTICE, build_document
from app.exporters.json_exporter import (
    JSON_SCHEMA_VERSION,
    build_json_document,
    validate_json_document,
)
from app.filename_template import DEFAULT_TEMPLATE
from tests.conftest import TURKISH_TEXT

TURKISH_CHARACTERS = "ğüşıİöçĞÜŞÖÇ"
ALL_FORMATS = ["md", "txt", "json", "docx", "pdf"]


def read_docx(path: Path) -> str:
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() for page in PdfReader(str(path)).pages)


def read_export(path: Path) -> str:
    if path.suffix == ".docx":
        return read_docx(path)
    if path.suffix == ".pdf":
        return read_pdf(path)
    return path.read_text(encoding="utf-8")


def test_registry_contains_all_formats():
    assert list(EXPORTERS) == ALL_FORMATS
    for format_id, exporter in EXPORTERS.items():
        assert exporter.format_id == format_id
        assert exporter.extension == format_id


def test_document_groups_consecutive_languages(session):
    document = build_document(session)
    assert [block.language for block in document.blocks] == ["en", "tr", "en"]
    assert dict(document.metadata)["Languages"] == "English, Turkish"
    assert dict(document.metadata)["Duration"] == "00:05:24"
    assert dict(document.metadata)["Date"] == "2026-10-03 14:30:00"


def test_markdown_export(session, tmp_path):
    path = export_session(session, tmp_path, DEFAULT_TEMPLATE, "md")
    text = path.read_text(encoding="utf-8")
    assert path.name == "2026-10-03_14-30-00_en-tr.md"
    assert text.startswith("# Speaking Session\n")
    assert "- Date: 2026-10-03 14:30:00" in text
    assert "- Languages: English, Turkish" in text
    assert "- Duration: 00:05:24" in text
    assert "## Transcript" in text
    assert text.count("### English") == 2
    assert text.count("### Turkish") == 1
    assert "[00:00:00] Yesterday I went to the gym." in text
    assert f"[00:00:02] {TURKISH_TEXT}" in text
    assert text.index("Yesterday") < text.index("Sonra") < text.index("We talked")


def test_text_export(session, tmp_path):
    text = export_session(session, tmp_path, DEFAULT_TEMPLATE, "txt").read_text(encoding="utf-8")
    assert text.startswith("Speaking Session\n================\n")
    assert "Languages: English, Turkish" in text
    assert "[English]\n[00:00:00] Yesterday I went to the gym." in text
    assert f"[Turkish]\n[00:00:02] {TURKISH_TEXT}" in text
    assert "#" not in text


def test_timestamps_can_be_disabled(session, tmp_path):
    options = ExportOptions(include_timestamps=False)
    text = export_session(session, tmp_path, DEFAULT_TEMPLATE, "txt", options).read_text("utf-8")
    assert "[00:00" not in text
    assert "\nYesterday I went to the gym.\n" in text


def test_json_export_matches_schema(session, tmp_path):
    path = export_session(session, tmp_path, DEFAULT_TEMPLATE, "json")
    document = json.loads(path.read_text(encoding="utf-8"))
    assert validate_json_document(document) == []
    assert document["schema_version"] == JSON_SCHEMA_VERSION
    assert document["session"]["started_at"] == "2026-10-03T14:30:00+03:00"
    assert document["session"]["duration_seconds"] == 324.0
    assert document["session"]["languages"] == ["en", "tr"]
    assert document["session"]["language_names"] == ["English", "Turkish"]
    assert [segment["language"] for segment in document["segments"]] == ["en", "tr", "en"]
    assert document["segments"][1] == {
        "index": 1,
        "start": 2.5,
        "end": 4.0,
        "language": "tr",
        "language_probability": 0.97,
        "text": TURKISH_TEXT,
    }
    # Non-ASCII text is stored readably, not as \u escapes.
    assert "arkadaşımı" in path.read_text(encoding="utf-8")


@pytest.mark.parametrize(
    "mutate, problem",
    [
        (lambda d: d.update(schema_version=99), "schema_version"),
        (lambda d: d.pop("session"), "session"),
        (lambda d: d["session"].pop("languages"), "languages"),
        (lambda d: d["session"].update(duration_seconds=-1), "negative"),
        (lambda d: d["segments"][0].update(text=""), "text"),
        (lambda d: d["segments"][0].update(start=9, end=1), "ends before"),
        (lambda d: d["segments"][1].update(index=5), "index"),
        (lambda d: d["segments"][0].pop("language"), "language"),
        (lambda d: d["segments"][0].update(language_probability=1.5), "0..1"),
    ],
)
def test_json_validation_detects_problems(session, mutate, problem):
    document = build_json_document(session)
    mutate(document)
    problems = validate_json_document(document)
    assert any(problem in item for item in problems), problems


def test_docx_export(session, tmp_path):
    path = export_session(session, tmp_path, DEFAULT_TEMPLATE, "docx")
    text = read_docx(path)
    assert "Speaking Session" in text
    assert "Languages: English, Turkish" in text
    assert f"[00:00:02] {TURKISH_TEXT}" in text
    for character in TURKISH_CHARACTERS:
        assert character in text or character not in TURKISH_TEXT

    from docx import Document

    properties = Document(str(path)).core_properties
    assert properties.author == ""
    assert properties.title == "Speaking Session"


def test_pdf_export_keeps_turkish_characters(session, tmp_path):
    path = export_session(session, tmp_path, DEFAULT_TEMPLATE, "pdf")
    assert path.read_bytes().startswith(b"%PDF")
    text = read_pdf(path)
    assert "Speaking Session" in text
    assert "English, Turkish" in text
    assert "Yesterday I went to the gym." in text
    # Every Turkish-specific letter must survive the round trip.
    for word in ("arkadaşımı", "gördüm", "Çığlık", "öğün", "şişe", "İstanbul", "ILIK", "ığdır"):
        assert word in text, f"{word!r} was corrupted in the PDF"


@pytest.mark.skipif(sys.platform != "win32", reason="relies on fonts shipped with Windows")
def test_pdf_mixes_scripts_in_one_document(session, tmp_path):
    from dataclasses import replace

    samples = {
        "de": "Ich habe gestern einen schönen Spaziergang gemacht, über die Straße.",
        "fr": "Ensuite, nous sommes allés au marché, ça coûte très cher.",
        "ru": "Вчера вечером мы долго гуляли по городу.",
        "el": "Καλημέρα σε όλους.",
        "zh": "昨天我去了健身房。",
    }
    session.segments = [
        replace(session.segments[0], language=code, text=text) for code, text in samples.items()
    ]
    session.languages = list(samples)
    text = read_pdf(export_session(session, tmp_path, DEFAULT_TEMPLATE, "pdf"))
    for sample in samples.values():
        assert sample in text


def test_pdf_embeds_a_unicode_font(session, tmp_path):
    from pypdf import PdfReader

    path = export_session(session, tmp_path, DEFAULT_TEMPLATE, "pdf")
    fonts = PdfReader(str(path)).pages[0]["/Resources"]["/Font"]
    subtypes = {fonts[name]["/Subtype"] for name in fonts}
    assert "/TrueType" in subtypes  # an embedded font, not only the built-in ones
    embedded = [fonts[name] for name in fonts if fonts[name]["/Subtype"] == "/TrueType"]
    assert all("/FontFile2" in font["/FontDescriptor"] for font in embedded)


@pytest.mark.parametrize("format_id", ALL_FORMATS)
def test_every_format_contains_the_same_content(session, tmp_path, format_id):
    text = read_export(export_session(session, tmp_path, DEFAULT_TEMPLATE, format_id))
    for segment in session.segments:
        assert segment.text in text
    assert "abc12345" in text  # session id
    if format_id == "json":
        assert '"languages": [\n      "en",\n      "tr"' in text
    else:
        assert "English, Turkish" in text
        assert "00:05:24" in text


@pytest.mark.parametrize("format_id", ALL_FORMATS)
def test_empty_transcript(empty_session, tmp_path, format_id):
    path = export_session(empty_session, tmp_path, DEFAULT_TEMPLATE, format_id)
    assert path.name == f"2026-10-03_09-00-00_none.{format_id}"
    text = read_export(path)
    if format_id == "json":
        document = json.loads(text)
        assert document["segments"] == []
        assert validate_json_document(document) == []
    else:
        assert EMPTY_NOTICE in text
        assert "None detected" in text


@pytest.mark.parametrize("format_id", ALL_FORMATS)
def test_special_characters_do_not_break_export(session, tmp_path, format_id):
    from dataclasses import replace

    tricky = 'Tags <b>&amp; "quotes" \x07 control, emoji-free: a < b > c & d'
    session.segments[0] = replace(session.segments[0], text=tricky)
    text = read_export(export_session(session, tmp_path, DEFAULT_TEMPLATE, format_id))
    assert "a < b > c & d" in text or format_id == "json"
    if format_id == "json":
        assert json.loads(text)["segments"][0]["text"] == tricky


def test_collision_adds_numeric_suffix(session, tmp_path):
    names = [export_session(session, tmp_path, "fixed", "txt").name for _ in range(3)]
    assert names == ["fixed.txt", "fixed_2.txt", "fixed_3.txt"]


def test_existing_file_is_never_overwritten(session, tmp_path):
    existing = tmp_path / "fixed.md"
    existing.write_text("precious", encoding="utf-8")
    path = export_session(session, tmp_path, "fixed", "md")
    assert path.name == "fixed_2.md"
    assert existing.read_text(encoding="utf-8") == "precious"


def test_missing_directory_is_created(session, tmp_path):
    target = tmp_path / "new" / "nested folder"
    path = export_session(session, target, DEFAULT_TEMPLATE, "md")
    assert path.parent == target and path.exists()


def test_no_temporary_files_remain(session, tmp_path):
    for format_id in ALL_FORMATS:
        export_session(session, tmp_path, DEFAULT_TEMPLATE, format_id)
    assert sorted(p.suffix for p in tmp_path.iterdir()) == sorted(f".{f}" for f in ALL_FORMATS)


def test_unknown_format(session, tmp_path):
    with pytest.raises(ExportError) as error:
        export_session(session, tmp_path, DEFAULT_TEMPLATE, "exe")
    assert error.value.code == "unknown_format"


def test_directory_that_is_a_file(session, tmp_path):
    blocker = tmp_path / "not-a-folder"
    blocker.write_text("x")
    with pytest.raises(ExportError) as error:
        export_session(session, blocker, DEFAULT_TEMPLATE, "md")
    assert error.value.code == "directory_invalid"


def test_check_directory_reports_missing(tmp_path):
    with pytest.raises(ExportError) as error:
        check_directory(tmp_path / "missing")
    assert error.value.code == "directory_missing"
    check_directory(tmp_path / "missing", create=True)
    assert (tmp_path / "missing").is_dir()
    assert list((tmp_path / "missing").iterdir()) == []  # the probe file is removed


def test_failed_export_leaves_no_file_and_reports_error(session, tmp_path, monkeypatch):
    def explode(self, session, path, options):
        path.write_text("partial")
        raise OSError("disk full")

    monkeypatch.setattr(type(EXPORTERS["md"]), "write", explode)
    with pytest.raises(ExportError) as error:
        export_session(session, tmp_path, DEFAULT_TEMPLATE, "md")
    assert error.value.code == "export_failed"
    assert "disk full" in error.value.detail
    assert list(tmp_path.iterdir()) == []


def test_exporter_that_writes_nothing_is_an_error(session, tmp_path, monkeypatch):
    monkeypatch.setattr(type(EXPORTERS["txt"]), "write", lambda *args: None)
    with pytest.raises(ExportError):
        export_session(session, tmp_path, DEFAULT_TEMPLATE, "txt")
    assert list(tmp_path.iterdir()) == []


def test_pdf_without_any_font_fails_cleanly(session, tmp_path, monkeypatch):
    from app.exporters import pdf_exporter

    monkeypatch.setattr(pdf_exporter, "_fonts", pdf_exporter._FontLibrary())
    monkeypatch.setattr(pdf_exporter, "font_candidates", lambda: [])
    with pytest.raises(ExportError) as error:
        export_session(session, tmp_path, DEFAULT_TEMPLATE, "pdf")
    assert error.value.code == "pdf_font_missing"
    assert list(tmp_path.iterdir()) == []


def test_export_to_explicit_path(session, tmp_path):
    target = tmp_path / "chosen name.json"
    assert export_to_path(session, target, "json") == target
    assert validate_json_document(json.loads(target.read_text(encoding="utf-8"))) == []


def test_failed_save_as_keeps_the_existing_file(session, tmp_path, monkeypatch):
    target = tmp_path / "existing.md"
    target.write_text("old content", encoding="utf-8")

    def explode(self, session, path, options):
        raise OSError("no access")

    monkeypatch.setattr(type(EXPORTERS["md"]), "write", explode)
    with pytest.raises(ExportError):
        export_to_path(session, target, "md")
    assert target.read_text(encoding="utf-8") == "old content"
