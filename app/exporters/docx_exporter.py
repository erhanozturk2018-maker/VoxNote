"""Microsoft Word export using python-docx."""

from __future__ import annotations

from pathlib import Path

from app.exporters.base import (
    EMPTY_NOTICE,
    TRANSCRIPT_HEADING,
    Exporter,
    ExportOptions,
    build_document,
    printable,
)
from app.transcript_models import Session


class DocxExporter(Exporter):
    format_id = "docx"
    extension = "docx"
    label = "Microsoft Word (.docx)"

    def write(self, session: Session, path: Path, options: ExportOptions) -> None:
        from docx import Document

        document = build_document(session, options)
        output = Document()

        # Do not leak library or user names into the file properties.
        properties = output.core_properties
        properties.title = document.title
        properties.author = ""
        properties.last_modified_by = ""
        properties.comments = ""
        properties.created = session.started_at.replace(tzinfo=None)

        if document.show_headings:
            output.add_heading(document.title, level=0)
        for name, value in document.metadata:
            paragraph = output.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{name}: ").bold = True
            paragraph.add_run(printable(value))

        if document.show_headings:
            output.add_heading(TRANSCRIPT_HEADING, level=1)
        if not document.blocks:
            output.add_paragraph().add_run(EMPTY_NOTICE).italic = True
        elif document.as_paragraph:
            output.add_paragraph(printable(document.paragraph))
        for block in () if document.as_paragraph else document.blocks:
            if document.show_language_headings:
                output.add_heading(block.language_name, level=2)
            for segment in block.segments:
                output.add_paragraph(printable(document.line(segment)))

        output.save(str(path))
